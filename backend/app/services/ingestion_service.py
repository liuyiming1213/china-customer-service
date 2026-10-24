import os
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from app.config import Settings
from app.repositories.sqlite_repository import SqliteRepository
from app.services.retrieval_service import ChromaKnowledgeService


def clean_text(value: object) -> str:
    return str(value or "").replace("\r", "").strip()


def build_case_text(thread_id: str, messages: list[tuple[int, str]]) -> tuple[str, dict]:
    ordered = sorted(messages, key=lambda item: item[0])
    full_text = "\n\n".join(f"消息 {branch}: {body}" for branch, body in ordered if body)
    if len(full_text) > 900:
        text = full_text[:600] + "\n……中间内容已省略……\n" + full_text[-300:]
    else:
        text = full_text
    first_message = ordered[0][1] if ordered else ""
    metadata = {
        "thread_id": thread_id,
        "source_quality": "historical",
        "allowed_for_auto_reply": False,
        "risk_level": "unknown",
        "message_count": len(ordered),
        "has_customer_reply": any("回答" in body or "ご返信" in body for _, body in ordered[1:]),
        "preview": first_message[:120],
    }
    return text, metadata


def ingest_workbook(settings: Settings, repository: SqliteRepository, knowledge: ChromaKnowledgeService) -> dict[str, int]:
    source = Path(settings.source_xlsx_path)
    if not source.exists():
        raise FileNotFoundError(f"Source workbook not found: {source}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    threads: dict[str, list[tuple[int, str]]] = defaultdict(list)
    skipped = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) < 3 or row[0] is None or row[1] is None:
            skipped += 1
            continue
        thread_id = str(row[0])
        try:
            branch = int(row[1])
        except (TypeError, ValueError):
            skipped += 1
            continue
        body = clean_text(row[2])
        if body:
            threads[thread_id].append((branch, body))
    max_threads = int(os.getenv("INGEST_MAX_THREADS", "0"))
    if max_threads > 0:
        threads = dict(list(threads.items())[:max_threads])
    documents: list[str] = []
    metadatas: list[dict] = []
    ids: list[str] = []
    for thread_id, messages in threads.items():
        document, metadata = build_case_text(thread_id, messages)
        if not document:
            continue
        documents.append(document)
        metadatas.append(metadata)
        ids.append(f"case-{thread_id}")
    indexed = knowledge.add_documents("support_cases", documents, metadatas, ids)
    repository.save_ingestion_run(str(source), sum(len(messages) for messages in threads.values()), len(threads), indexed)
    return {
        "imported_threads": len(threads),
        "imported_messages": sum(len(messages) for messages in threads.values()),
        "indexed_cases": indexed,
        "skipped_rows": skipped,
    }
