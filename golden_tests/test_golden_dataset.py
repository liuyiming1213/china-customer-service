from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
GOLDEN_FILE = ROOT / "黄金测试集_脱敏客服.xlsx"
EXPECTED_DIMENSIONS = {
    "回复生成",
    "风险判断",
    "知识检索",
    "SSE可用性",
    "外部模型异常",
    "输入安全",
    "记忆写入",
    "知识库导入",
    "基础接口",
}


def load_golden_rows():
    workbook = openpyxl.load_workbook(GOLDEN_FILE, read_only=True, data_only=True)
    sheet = workbook["黄金测试集"]
    rows = list(sheet.iter_rows(min_row=5, values_only=True))
    return workbook, sheet, rows


def test_golden_workbook_contract():
    assert GOLDEN_FILE.exists()
    workbook, sheet, rows = load_golden_rows()
    assert set(workbook.sheetnames) == {"黄金测试集", "覆盖矩阵", "抽样说明"}
    assert len(rows) == 90
    assert all(row[0].startswith("GOLD-") for row in rows)
    assert len({row[0] for row in rows}) == 90
    assert len({row[1] for row in rows}) == 90


def test_golden_records_cover_boundaries_and_risk_values():
    _, _, rows = load_golden_rows()
    risks = {row[5] for row in rows}
    tags = "、".join(str(row[9] or "") for row in rows)
    assert risks <= {"low", "medium", "high"}
    assert {"low", "medium", "high"} <= risks
    assert "空输入边界" in tags
    assert "超长输入" in tags
    assert "字段缺失" in tags
    assert "多轮对话" in tags


def test_golden_coverage_matrix_contains_all_required_dimensions():
    workbook = openpyxl.load_workbook(GOLDEN_FILE, read_only=True, data_only=True)
    sheet = workbook["覆盖矩阵"]
    dimensions = {row[0] for row in sheet.iter_rows(min_row=5, max_col=1, values_only=True) if row[0]}
    assert EXPECTED_DIMENSIONS <= dimensions
